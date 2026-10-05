#!/usr/bin/env python3

import argparse
import html
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from jsonschema import Draft202012Validator
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


NAVY = "1F3864"
BLUE = "2E75B6"
LIGHT_BLUE = "EBF3FB"
LIGHT_YELLOW = "FFF2CC"
HEADER_FILL = "D9E1F2"
PASS_FILL = "E2EFDA"
BLOCKED_FILL = "FCE4D6"
FAIL_FILL = "F4CCCC"
NOT_RUN_FILL = "E7E6E6"
ZEBRA_FILL = "F7F9FC"
RISK_COLORS = {
    "HIGH": "E53935",
    "ELEVATED": "FB8C00",
    "MEDIUM": "FFA726",
    "LOW": "43A047",
}
RISK_RANK = {"LOW": 1, "MEDIUM": 2, "ELEVATED": 3, "HIGH": 4}


def parse_args():
    parser = argparse.ArgumentParser(
        description="Build Netra DOCX, XLSX, and HTML artifacts from analysis.json."
    )
    parser.add_argument("analysis", type=Path)
    parser.add_argument("--output-dir", type=Path)
    return parser.parse_args()


def validate_analysis(schema_path, analysis):
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    errors = sorted(
        Draft202012Validator(schema).iter_errors(analysis),
        key=lambda error: list(error.absolute_path),
    )
    messages = []
    for error in errors:
        location = "/" + "/".join(str(part) for part in error.absolute_path)
        messages.append(f"{location or '/'} {error.message}")
    if not errors:
        messages.extend(release_summary_errors(analysis))
    if messages:
        raise SystemExit("Invalid Netra analysis:\n" + "\n".join(messages))


def release_summary_errors(analysis):
    tickets = analysis["tickets"]
    summary = analysis["releaseSummary"]
    verdict_counts = {
        verdict: sum(ticket["readiness"]["verdict"] == verdict for ticket in tickets)
        for verdict in ("ready", "warning", "blocked")
    }
    total_test_cases = sum(
        len(ticket["testCoverage"]["testCases"]) for ticket in tickets
    )
    expected = {
        "totalTickets": len(tickets),
        "ready": verdict_counts["ready"],
        "warnings": verdict_counts["warning"],
        "blocked": verdict_counts["blocked"],
        "totalTestCases": total_test_cases,
        "averageTestCasesPerTicket": round(total_test_cases / len(tickets), 2),
    }
    messages = [
        f"/releaseSummary/{field} must be {value!r}, got {summary[field]!r}"
        for field, value in expected.items()
        if summary[field] != value
    ]
    coverage_counts = {
        ticket["key"]: len(ticket["testCoverage"]["testCases"])
        for ticket in tickets
    }
    highest_count = max(coverage_counts.values())
    highest_keys = {
        key for key, count in coverage_counts.items() if count == highest_count
    }
    if summary["highestCoverageTicket"] not in highest_keys:
        messages.append(
            "/releaseSummary/highestCoverageTicket must identify a ticket with "
            f"{highest_count} test case(s)"
        )
    return messages


def safe_identifier(value):
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._-")
    return value or "netra"


def add_field(paragraph, instruction, placeholder=None):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    field = OxmlElement("w:instrText")
    field.set(qn("xml:space"), "preserve")
    field.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    elements = [begin, field, separate]
    if placeholder:
        result = OxmlElement("w:t")
        result.text = placeholder
        elements.append(result)
    elements.append(end)
    run._r.extend(elements)


def shade_paragraph(paragraph, color):
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def shade_cell(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def style_docx_header(cells):
    for cell in cells:
        shade_cell(cell, HEADER_FILL)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(NAVY)


def set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def prevent_row_split(row):
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))


def set_column_widths(table, widths):
    table.autofit = False
    for grid_column, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_column.set(qn("w:w"), str(Inches(width).twips))
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_bullets(document, items, empty_text="None"):
    if not items:
        document.add_paragraph(empty_text)
        return
    for item in items:
        document.add_paragraph(str(item), style="List Bullet")


def source_label(source):
    label = f"{source['type']}: {source['id']}"
    if source.get("url"):
        label += f" ({source['url']})"
    return label


def configure_docx(document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 13)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(
            NAVY if style_name in {"Title", "Heading 1"} else BLUE
        )
        borders = style.element.pPr.find(qn("w:pBdr")) if style.element.pPr is not None else None
        if borders is not None:
            style.element.pPr.remove(borders)

    if "Summary Label" not in styles:
        style = styles.add_style("Summary Label", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.font.bold = True

    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")


def add_release_summary(document, summary):
    document.add_heading("Release Summary", level=1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Total tickets", summary["totalTickets"]),
        ("Ready", summary["ready"]),
        ("Warnings", summary["warnings"]),
        ("Blocked", summary["blocked"]),
        ("Total test cases", summary["totalTestCases"]),
        ("Average test cases per ticket", summary["averageTestCasesPerTicket"]),
        ("Highest coverage ticket", summary["highestCoverageTicket"]),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)
        cells[0].paragraphs[0].runs[0].bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    for heading, values in (
        ("Live Dependencies", summary["liveDependencies"]),
        ("Production Risk Themes", summary["productionRiskThemes"]),
        ("Open Ambiguities", summary["openAmbiguities"]),
    ):
        document.add_heading(heading, level=2)
        add_bullets(document, values)


def build_docx(analysis, output_path):
    document = Document()
    configure_docx(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Netra Deep Analysis")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(analysis["input"]["value"]).italic = True
    generated = document.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated.add_run(f"Generated {analysis['createdAt']}")
    document.add_page_break()

    add_release_summary(document, analysis["releaseSummary"])
    document.add_heading("Contents", level=1)
    contents = document.add_paragraph()
    add_field(
        contents,
        'TOC \\o "1-2" \\h \\z \\u',
        "Open in Word and update this field to refresh the table of contents.",
    )

    for ticket in analysis["tickets"]:
        document.add_page_break()
        document.add_heading(f"{ticket['key']} {ticket['summary']}", level=1)
        metadata = ticket.get("metadata") or {}
        metadata_table = document.add_table(rows=2, cols=7)
        metadata_table.style = "Table Grid"
        labels = [
            "Type",
            "Status",
            "Priority",
            "Assignee",
            "Components",
            "Fix Version",
            "Triage",
        ]
        values = [
            metadata.get("issueType") or "Not specified",
            metadata.get("status") or "Not specified",
            metadata.get("priority") or "Not specified",
            metadata.get("assignee") or "Not specified",
            ", ".join(metadata.get("components") or []) or "Not specified",
            ", ".join(metadata.get("fixVersions") or []) or "Not specified",
            ticket["readiness"]["verdict"].upper(),
        ]
        for cell, label in zip(metadata_table.rows[0].cells, labels):
            cell.text = label
        style_docx_header(metadata_table.rows[0].cells)
        set_repeat_table_header(metadata_table.rows[0])
        set_column_widths(metadata_table, [0.70, 0.80, 0.75, 1.15, 1.25, 1.15, 0.90])
        for cell, value in zip(metadata_table.rows[1].cells, values):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        verdict_fill = {
            "ready": "43A047",
            "warning": "FFA726",
            "blocked": "E53935",
        }[ticket["readiness"]["verdict"]]
        shade_cell(metadata_table.rows[1].cells[-1], verdict_fill)
        if ticket.get("contextNote"):
            note = document.add_paragraph(ticket["contextNote"])
            shade_paragraph(note, LIGHT_BLUE)

        sections = [
            ("Background", ticket["background"]),
            ("Intent", ticket["intent"]),
            ("Cross Functional Impact", ticket["crossFunctionalImpact"]),
            ("Requirement Gaps", [item["description"] for item in ticket["gaps"]]),
            ("Ambiguities", [item["description"] for item in ticket["ambiguities"]]),
            (
                "Dependencies",
                [
                    f"{item['id']} | {item['status']} | "
                    f"{'blocking' if item['blocking'] else 'non-blocking'}"
                    for item in ticket["dependencies"]
                ],
            ),
        ]
        for heading, content in sections:
            document.add_heading(heading, level=2)
            if isinstance(content, list):
                add_bullets(document, content)
            else:
                document.add_paragraph(content)

            if heading == "Dependencies" and ticket["navigationContext"]["frameworkGap"]:
                navigation = ticket["navigationContext"]
                module_label = (
                    f" (module {navigation['module']})" if navigation["module"] else ""
                )
                note = document.add_paragraph(
                    "Automation heads-up: Navigation node "
                    f"{navigation['node'] or 'Unmapped'}"
                    f"{module_label} "
                    "is marked as a framework gap. Downstream automation may require "
                    "navigation or page-model support."
                )
                shade_paragraph(note, LIGHT_BLUE)

        document.add_heading("Historical Analysis", level=2)
        risks = ticket["productionRisks"]
        if risks:
            risk_table = document.add_table(rows=1, cols=4)
            risk_table.style = "Table Grid"
            headers = ["Bug", "Summary", "Risk", "Why it matters here"]
            for cell, label in zip(risk_table.rows[0].cells, headers):
                cell.text = label
            style_docx_header(risk_table.rows[0].cells)
            set_repeat_table_header(risk_table.rows[0])
            set_column_widths(risk_table, [0.85, 1.45, 1.05, 3.35])
            for risk in risks:
                row = risk_table.add_row()
                prevent_row_split(row)
                cells = row.cells
                values = [
                    risk["issueKey"],
                    risk.get("summary", "Not specified"),
                    risk["rating"],
                    risk["reason"],
                ]
                for cell, value in zip(cells, values):
                    cell.text = value
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                shade_cell(cells[2], RISK_COLORS[risk["rating"]])
            overall = max(risks, key=lambda item: RISK_RANK[item["rating"]])["rating"]
            paragraph = document.add_paragraph()
            run = paragraph.add_run(f"Overall historical risk: {overall}")
            run.bold = True
            run.font.color.rgb = RGBColor.from_string(RISK_COLORS[overall])
        else:
            document.add_paragraph("None")

        document.add_heading("References", level=2)
        add_bullets(document, [source_label(item) for item in ticket["sourceRefs"]])

        blocking = [item for item in ticket["dependencies"] if item["blocking"]]
        if blocking:
            note = document.add_paragraph(
                "Blocking dependencies: " + ", ".join(item["id"] for item in blocking)
            )
            shade_paragraph(note, LIGHT_YELLOW)

    document.save(output_path)


def unique_sheet_name(key, used):
    base = re.sub(r"[\\/*?:\[\]]", "_", key)[:31] or "Ticket"
    name = base
    suffix = 2
    while name in used:
        tail = f"_{suffix}"
        name = f"{base[:31 - len(tail)]}{tail}"
        suffix += 1
    used.add(name)
    return name


def style_header(row):
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Arial", size=10, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_xlsx(analysis, output_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()
    headers = [
        "Test ID",
        "Requirement ID",
        "Test Description",
        "Test Data",
        "Expected Result",
        "Actual Result",
        "Test Status",
        "Owner",
        "Comments",
    ]
    widths = [18, 20, 55, 40, 50, 20, 16, 18, 40]

    for ticket in analysis["tickets"]:
        sheet = workbook.create_sheet(unique_sheet_name(ticket["key"], used_names))
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.append(headers)
        style_header(sheet[1])

        cases = ticket["testCoverage"]["testCases"]
        for index, case in enumerate(cases, start=2):
            display_status = {
                "not_executed": "NOT RUN",
                "blocked": "BLOCKED",
                "pass": "PASS",
                "fail": "FAIL",
            }[case["status"]]
            requirement_labels = [item.rsplit("-", 1)[-1] for item in case["requirementIds"]]
            type_label = case["type"].upper()
            qualifier = f"{type_label} — {', '.join(requirement_labels)}"
            prefix = f"[{qualifier}]"
            description = case["description"]
            if not description.startswith(prefix):
                description = f"{prefix} {description.rstrip('.')}"
            sheet.append(
                [
                    case["id"],
                    ", ".join(case["requirementIds"]),
                    description,
                    case["testData"],
                    case["expectedResult"],
                    case["actualResult"],
                    display_status,
                    case["owner"],
                    case["comments"],
                ]
            )
            fill = None
            if display_status == "PASS":
                fill = PatternFill("solid", fgColor=PASS_FILL)
            elif display_status == "BLOCKED":
                fill = PatternFill("solid", fgColor=BLOCKED_FILL)
            elif display_status == "FAIL":
                fill = PatternFill("solid", fgColor=FAIL_FILL)
            elif display_status == "NOT RUN":
                fill = PatternFill("solid", fgColor=NOT_RUN_FILL)
            elif index % 2 == 0:
                fill = PatternFill("solid", fgColor=ZEBRA_FILL)
            for cell in sheet[index]:
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill
            sheet.row_dimensions[index].height = 35

        sheet.auto_filter.ref = f"A1:I{max(sheet.max_row, 1)}"
        verdict = ticket["readiness"]["verdict"]
        sheet.sheet_properties.tabColor = {
            "ready": "43A047",
            "warning": "FFA726",
            "blocked": "E53935",
        }[verdict]
        for column, width in zip("ABCDEFGHI", widths):
            sheet.column_dimensions[column].width = width

    summary = workbook.create_sheet("Summary")
    summary.sheet_view.showGridLines = False
    summary_headers = [
        "Ticket Key",
        "Summary",
        "Triage Verdict",
        "Total TCs",
        "R (context items)",
        "Blockers",
        "Warnings",
    ]
    summary.append(summary_headers)
    style_header(summary[1])
    for ticket in analysis["tickets"]:
        verdict = ticket["readiness"]["verdict"]
        summary.append(
            [
                ticket["key"],
                ticket["summary"],
                verdict.upper(),
                len(ticket["testCoverage"]["testCases"]),
                ticket["testCoverage"]["requirementCount"],
                1 if verdict == "blocked" else 0,
                1 if verdict == "warning" else 0,
            ]
        )
    first_ticket_row = 2
    last_ticket_row = 1 + len(analysis["tickets"])
    summary.append(
        [
            "TOTAL",
            f"{len(analysis['tickets'])} ticket(s)",
            "",
            f"=SUM(D{first_ticket_row}:D{last_ticket_row})",
            f"=SUM(E{first_ticket_row}:E{last_ticket_row})",
            f"=SUM(F{first_ticket_row}:F{last_ticket_row})",
            f"=SUM(G{first_ticket_row}:G{last_ticket_row})",
        ]
    )
    for cell in summary[summary.max_row]:
        cell.font = Font(name="Arial", size=10, bold=True)
    summary.append([])
    summary.append(["Status legend"])
    legend = [
        ("PASS", "Executed and passed", PASS_FILL),
        ("FAIL", "Executed and failed", FAIL_FILL),
        ("BLOCKED", "Cannot be executed", BLOCKED_FILL),
        ("NOT RUN", "Designed, awaiting execution", NOT_RUN_FILL),
        ("READY tab", "Triage: no flags", "43A047"),
        ("WARNING tab", "Triage: warnings", "FFA726"),
        ("BLOCKED tab", "Triage: blocker", "E53935"),
    ]
    for label, description, color in legend:
        summary.append([label, description])
        summary.cell(summary.max_row, 1).fill = PatternFill("solid", fgColor=color)
    for row in summary.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, bold=cell.font.bold)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column, width in zip("ABCDEFG", [16, 50, 18, 12, 18, 12, 12]):
        summary.column_dimensions[column].width = width
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = f"A1:G{1 + len(analysis['tickets'])}"
    workbook.save(output_path)


def list_html(items):
    if not items:
        return "<p>None</p>"
    return "<ul>" + "".join(f"<li>{html.escape(str(item))}</li>" for item in items) + "</ul>"


def build_html(analysis, output_path):
    summary = analysis["releaseSummary"]
    total = max(summary["totalTickets"], 1)
    ready_width = summary["ready"] * 100 / total
    warning_width = summary["warnings"] * 100 / total
    blocked_width = summary["blocked"] * 100 / total
    highest_ticket = next(
        (
            ticket
            for ticket in analysis["tickets"]
            if ticket["key"] == summary["highestCoverageTicket"]
        ),
        None,
    )
    highest_detail = html.escape(summary["highestCoverageTicket"])
    if highest_ticket:
        highest_detail += (
            f"<small>{len(highest_ticket['testCoverage']['testCases'])} TCs · "
            f"R={highest_ticket['testCoverage']['requirementCount']}</small>"
        )
    ticket_rows = "".join(
        "<tr>"
        f"<td>{html.escape(ticket['key'])}</td>"
        f"<td>{html.escape(ticket['summary'])}</td>"
        f"<td><span class=\"pill {html.escape(ticket['readiness']['verdict'])}\">"
        f"{html.escape(ticket['readiness']['verdict'].upper())}</span></td>"
        f"<td>{len(ticket['testCoverage']['testCases'])}</td>"
        f"<td>{ticket['testCoverage']['requirementCount']}</td>"
        f"<td>{html.escape(ticket['navigationContext']['node'] or 'Unmapped')}</td>"
        "</tr>"
        for ticket in analysis["tickets"]
    )
    all_risks = [
        (ticket["key"], risk)
        for ticket in analysis["tickets"]
        for risk in ticket["productionRisks"]
    ]
    risk_rows = "".join(
        "<tr>"
        f"<td>{html.escape(ticket_key)}</td>"
        f"<td>{html.escape(risk['issueKey'])}</td>"
        f"<td>{html.escape(risk.get('summary', 'Not specified'))}</td>"
        f"<td><span class=\"pill risk-{risk['rating'].lower()}\">"
        f"{html.escape(risk['rating'])}</span></td>"
        f"<td>{html.escape(risk['reason'])}</td>"
        "</tr>"
        for ticket_key, risk in all_risks
    )
    content = f"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Netra Release Readiness</title>
  <style>
    :root {{ --navy:#1F3864; --blue:#2E75B6; --green:#43A047; --amber:#FFA726; --red:#E53935; }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; background:#f5f7fa; color:#1f2937; font-family:Arial,sans-serif; }}
    main {{ max-width:1180px; margin:0 auto; padding:32px 20px; }}
    h1,h2 {{ color:var(--navy); }}
    h1 {{ margin:0 0 6px; }}
    .meta {{ color:#5f6b7a; margin-bottom:24px; }}
    .grid {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(190px,1fr)); gap:16px; }}
    .card {{ background:#fff; border-radius:8px; padding:18px; box-shadow:0 2px 8px rgba(31,56,100,.10); }}
    .label {{ color:#5f6b7a; font-size:13px; }}
    .value {{ color:var(--navy); font-size:30px; font-weight:700; margin-top:6px; }}
    .value small {{ display:block; color:#5f6b7a; font-size:12px; font-weight:400; margin-top:4px; }}
    .bar {{ display:flex; overflow:hidden; height:24px; border-radius:6px; background:#e5e7eb; }}
    .bar .ready {{ width:{ready_width:.2f}%; background:var(--green); }}
    .bar .warning {{ width:{warning_width:.2f}%; background:var(--amber); }}
    .bar .blocked {{ width:{blocked_width:.2f}%; background:var(--red); }}
    section {{ margin-top:24px; }}
    ul {{ padding-left:20px; }}
    .table-wrap {{ overflow-x:auto; }}
    table {{ width:100%; border-collapse:collapse; font-size:14px; }}
    th {{ background:#D9E1F2; color:var(--navy); text-align:left; padding:9px 10px; }}
    td {{ border-top:1px solid #e5e7eb; padding:9px 10px; vertical-align:top; }}
    .pill {{ display:inline-block; border-radius:999px; color:#fff; font-size:12px; font-weight:700; padding:2px 8px; }}
    .pill.ready,.risk-low {{ background:var(--green); }}
    .pill.warning,.risk-medium,.risk-elevated {{ background:var(--amber); }}
    .pill.blocked,.risk-high {{ background:var(--red); }}
  </style>
</head>
<body>
<main>
  <h1>Release Readiness</h1>
  <div class="meta">{html.escape(analysis['input']['value'])} | Generated {html.escape(analysis['createdAt'])}</div>
  <div class="grid">
    <div class="card"><div class="label">Total tickets</div><div class="value">{summary['totalTickets']}</div></div>
    <div class="card"><div class="label">Total test cases</div><div class="value">{summary['totalTestCases']}</div></div>
    <div class="card"><div class="label">Average test cases per ticket</div><div class="value">{summary['averageTestCasesPerTicket']}</div></div>
    <div class="card"><div class="label">Highest coverage ticket</div><div class="value">{highest_detail}</div></div>
  </div>
  <section class="card">
    <h2>Readiness</h2>
    <div class="bar" aria-label="Ready {summary['ready']}, warnings {summary['warnings']}, blocked {summary['blocked']}">
      <div class="ready"></div><div class="warning"></div><div class="blocked"></div>
    </div>
    <p>Ready {summary['ready']} | Warnings {summary['warnings']} | Blocked {summary['blocked']}</p>
  </section>
  <section class="card">
    <h2>Tickets</h2>
    <div class="table-wrap"><table>
      <thead><tr><th>Ticket</th><th>Summary</th><th>Triage</th><th>TCs</th><th>R</th><th>Navigation node</th></tr></thead>
      <tbody>{ticket_rows}</tbody>
    </table></div>
  </section>
  <section class="card">
    <h2>Production Risk Details</h2>
    <div class="table-wrap"><table>
      <thead><tr><th>Ticket</th><th>Bug</th><th>Summary</th><th>Risk</th><th>Why it matters here</th></tr></thead>
      <tbody>{risk_rows or '<tr><td colspan="5">None</td></tr>'}</tbody>
    </table></div>
  </section>
  <section class="grid">
    <div class="card"><h2>Production Risk Themes</h2>{list_html(summary['productionRiskThemes'])}</div>
    <div class="card"><h2>Live Dependencies</h2>{list_html(summary['liveDependencies'])}</div>
    <div class="card"><h2>Open Ambiguities</h2>{list_html(summary['openAmbiguities'])}</div>
  </section>
</main>
</body>
</html>
"""
    output_path.write_text(content, encoding="utf-8")


def verify_outputs(analysis, docx_path, xlsx_path, html_path):
    for path in (docx_path, xlsx_path, html_path):
        if not path.is_file() or path.stat().st_size == 0:
            raise RuntimeError(f"Missing or empty artifact: {path}")
    if not zipfile.is_zipfile(docx_path) or not zipfile.is_zipfile(xlsx_path):
        raise RuntimeError("DOCX or XLSX output is not a valid OOXML archive")

    document = Document(docx_path)
    heading_text = {paragraph.text for paragraph in document.paragraphs}
    for ticket in analysis["tickets"]:
        expected = f"{ticket['key']} {ticket['summary']}"
        if expected not in heading_text:
            raise RuntimeError(f"DOCX is missing ticket section: {ticket['key']}")

    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    if "Summary" not in workbook.sheetnames:
        raise RuntimeError("XLSX is missing Summary")
    ticket_sheets = [name for name in workbook.sheetnames if name != "Summary"]
    if len(ticket_sheets) != len(analysis["tickets"]):
        raise RuntimeError("XLSX ticket sheet count does not match analysis.json")
    for sheet_name, ticket in zip(ticket_sheets, analysis["tickets"]):
        rows = workbook[sheet_name].max_row - 1
        if rows != len(ticket["testCoverage"]["testCases"]):
            raise RuntimeError(f"XLSX row count mismatch for {ticket['key']}")

    html_text = html_path.read_text(encoding="utf-8")
    if "Release Readiness" not in html_text or "<html" not in html_text:
        raise RuntimeError("HTML dashboard is incomplete")


def main():
    args = parse_args()
    skill_root = Path(__file__).resolve().parent.parent
    repo_root = skill_root.parents[2]
    schema_path = repo_root / ".claude/contracts/analysis-output.schema.json"
    analysis_path = args.analysis.resolve()
    analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
    validate_analysis(schema_path, analysis)
    output_dir = (args.output_dir or repo_root / "target/netra").resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    identifier = safe_identifier(analysis["input"]["value"])
    date = analysis["createdAt"][:10]
    docx_path = output_dir / f"Deep_Analysis_{identifier}_{date}.docx"
    xlsx_path = output_dir / f"Test_Design_{identifier}_{date}.xlsx"
    html_path = output_dir / f"Release_Readiness_Dashboard_{identifier}_{date}.html"

    build_docx(analysis, docx_path)
    build_xlsx(analysis, xlsx_path)
    build_html(analysis, html_path)
    verify_outputs(analysis, docx_path, xlsx_path, html_path)

    analysis["compatibilityArtifacts"].update(
        {
            "deepAnalysisDocx": str(docx_path),
            "testDesignXlsx": str(xlsx_path),
            "readinessDashboardHtml": str(html_path),
        }
    )
    validate_analysis(schema_path, analysis)
    analysis_path.write_text(json.dumps(analysis, indent=2) + "\n", encoding="utf-8")

    for path in (docx_path, xlsx_path, html_path):
        print(path)


if __name__ == "__main__":
    main()

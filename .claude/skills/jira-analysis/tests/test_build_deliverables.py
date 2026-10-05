import copy
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook


TEST_ROOT = Path(__file__).resolve().parent
SKILL_ROOT = TEST_ROOT.parent
SCRIPT = SKILL_ROOT / "scripts/build-deliverables.py"
FIXTURES = TEST_ROOT / "fixtures"
PASS_FILL = "E2EFDA"
BLOCKED_FILL = "FCE4D6"
FAIL_FILL = "F4CCCC"
NOT_RUN_FILL = "E7E6E6"


class BuildDeliverablesTest(unittest.TestCase):
    def run_builder(self, fixture, output_dir):
        analysis_path = Path(fixture)
        if not analysis_path.is_absolute():
            analysis_path = FIXTURES / analysis_path
        return subprocess.run(
            [
                sys.executable,
                str(SCRIPT),
                str(analysis_path),
                "--output-dir",
                str(output_dir),
            ],
            capture_output=True,
            text=True,
            check=False,
        )

    def write_analysis(self, directory, analysis):
        directory = Path(directory)
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / "analysis.json"
        path.write_text(json.dumps(analysis), encoding="utf-8")
        return path

    def load_valid_analysis(self):
        return json.loads((FIXTURES / "valid-analysis.json").read_text(encoding="utf-8"))

    def test_valid_analysis_builds_all_artifacts(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            analysis_path = self.write_analysis(
                temporary_directory, self.load_valid_analysis()
            )
            output_dir = Path(temporary_directory) / "output"
            result = self.run_builder(analysis_path, output_dir)

            self.assertEqual(0, result.returncode, result.stderr)
            generated = {}
            for extension in ("docx", "xlsx", "html"):
                artifacts = list(output_dir.glob(f"*.{extension}"))
                self.assertEqual(1, len(artifacts))
                self.assertGreater(artifacts[0].stat().st_size, 0)
                generated[extension] = str(artifacts[0].resolve())

            updated_analysis = json.loads(analysis_path.read_text(encoding="utf-8"))
            self.assertEqual(
                {
                    "deepAnalysisDocx": generated["docx"],
                    "testDesignXlsx": generated["xlsx"],
                    "readinessDashboardHtml": generated["html"],
                    "legacyCacheJson": None,
                },
                updated_analysis["compatibilityArtifacts"],
            )

    def test_builder_preserves_human_readable_compatibility_features(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            analysis_path = self.write_analysis(
                temporary_directory, self.load_valid_analysis()
            )
            output_dir = Path(temporary_directory) / "output"
            result = self.run_builder(analysis_path, output_dir)
            self.assertEqual(0, result.returncode, result.stderr)

            workbook = load_workbook(next(output_dir.glob("*.xlsx")), data_only=False)
            ticket_sheet = workbook["QA-100"]
            self.assertEqual("NOT RUN", ticket_sheet["G2"].value)
            self.assertTrue(ticket_sheet["C2"].value.startswith("[POSITIVE — R01]"))
            self.assertEqual(NOT_RUN_FILL, ticket_sheet["G2"].fill.fgColor.rgb[-6:])
            summary = workbook["Summary"]
            self.assertEqual("=SUM(D2:D2)", summary["D3"].value)
            legend = {summary.cell(row, 1).value for row in range(1, summary.max_row + 1)}
            self.assertTrue({"PASS", "FAIL", "BLOCKED", "NOT RUN"}.issubset(legend))

            document = Document(next(output_dir.glob("*.docx")))
            table_headers = [tuple(cell.text for cell in table.rows[0].cells) for table in document.tables]
            self.assertIn(
                ("Type", "Status", "Priority", "Assignee", "Components", "Fix Version", "Triage"),
                table_headers,
            )
            self.assertIn(
                ("Bug", "Summary", "Risk", "Why it matters here"),
                table_headers,
            )
            self.assertIn(
                "Overall historical risk: HIGH",
                {paragraph.text for paragraph in document.paragraphs},
            )
            self.assertTrue(
                any(
                    paragraph.text.startswith("Automation heads-up:")
                    for paragraph in document.paragraphs
                )
            )

            dashboard = next(output_dir.glob("*.html")).read_text(encoding="utf-8")
            self.assertIn("<h2>Tickets</h2>", dashboard)
            self.assertIn("<h2>Production Risk Details</h2>", dashboard)
            self.assertIn("HT-100", dashboard)
            self.assertIn("<th>Ticket</th><th>Bug</th>", dashboard)

            with zipfile.ZipFile(next(output_dir.glob("*.docx"))) as archive:
                document_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertIn('TOC \\o "1-2" \\h \\z \\u', document_xml)
            self.assertGreaterEqual(document_xml.count("<w:tblHeader"), 2)
            self.assertIn("<w:cantSplit", document_xml)
            self.assertIn('<w:gridCol w:w="1512"', document_xml)

    def test_builder_accepts_pre_metadata_analysis(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            analysis = self.load_valid_analysis()
            analysis["tickets"][0].pop("metadata")
            analysis["tickets"][0]["productionRisks"][0].pop("summary")
            analysis_path = self.write_analysis(temporary_directory, analysis)
            output_dir = Path(temporary_directory) / "output"

            result = self.run_builder(analysis_path, output_dir)

            self.assertEqual(0, result.returncode, result.stderr)
            document = Document(next(output_dir.glob("*.docx")))
            table_text = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            self.assertIn("Not specified", table_text)
            dashboard = next(output_dir.glob("*.html")).read_text(encoding="utf-8")
            self.assertIn("Not specified", dashboard)

    def test_builder_handles_multiple_tickets_and_summary_verdicts(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            analysis = self.load_valid_analysis()
            warning_ticket = copy.deepcopy(analysis["tickets"][0])
            warning_ticket["key"] = "QA-101"
            warning_ticket["summary"] = "Warning ticket"
            warning_ticket["readiness"] = {
                "verdict": "warning",
                "flags": ["MISSING_ASSIGNEE", "NO_LINKED_CONTEXT"],
            }
            blocked_ticket = copy.deepcopy(analysis["tickets"][0])
            blocked_ticket["key"] = "QA-102"
            blocked_ticket["summary"] = "Blocked ticket"
            blocked_ticket["readiness"] = {
                "verdict": "blocked",
                "flags": ["MISSING_COMPONENTS"],
            }
            for ticket, number in ((warning_ticket, "101"), (blocked_ticket, "102")):
                ticket["requirements"][0]["id"] = f"QA-{number}-R01"
                ticket["sourceManifest"][0]["id"] = f"QA-{number}"
                for index, case in enumerate(ticket["testCoverage"]["testCases"], start=1):
                    case["id"] = f"TC_QA-{number}_{index:02d}"
                    case["requirementIds"] = [f"QA-{number}-R01"]
            analysis["tickets"].extend([warning_ticket, blocked_ticket])
            analysis["releaseSummary"].update(
                {
                    "totalTickets": 3,
                    "ready": 1,
                    "warnings": 1,
                    "blocked": 1,
                    "totalTestCases": 9,
                    "averageTestCasesPerTicket": 3,
                }
            )
            analysis_path = self.write_analysis(temporary_directory, analysis)
            output_dir = Path(temporary_directory) / "output"

            result = self.run_builder(analysis_path, output_dir)

            self.assertEqual(0, result.returncode, result.stderr)
            workbook = load_workbook(next(output_dir.glob("*.xlsx")), data_only=False)
            summary = workbook["Summary"]
            self.assertEqual([0, 1, 0], [summary[f"G{row}"].value for row in range(2, 5)])
            self.assertEqual([0, 0, 1], [summary[f"F{row}"].value for row in range(2, 5)])
            dashboard = next(output_dir.glob("*.html")).read_text(encoding="utf-8")
            for ticket_key in ("QA-100", "QA-101", "QA-102"):
                self.assertIn(f"<td>{ticket_key}</td><td>HT-100</td>", dashboard)

    def test_status_fills_and_existing_description_prefix_are_preserved(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            analysis = self.load_valid_analysis()
            cases = analysis["tickets"][0]["testCoverage"]["testCases"]
            cases[0]["description"] = "[POSITIVE — R01] Already prefixed description"
            cases[0]["status"] = "pass"
            cases[1]["status"] = "fail"
            blocked_case = copy.deepcopy(cases[1])
            blocked_case.update(
                {"id": "TC_QA-100_04", "type": "blocked", "status": "blocked"}
            )
            not_run_case = copy.deepcopy(cases[1])
            not_run_case.update({"id": "TC_QA-100_05", "status": "not_executed"})
            cases.extend([blocked_case, not_run_case])
            analysis["tickets"][0]["testCoverage"]["minimumTestCases"] = 5
            analysis["releaseSummary"].update(
                {"totalTestCases": 5, "averageTestCasesPerTicket": 5}
            )
            analysis_path = self.write_analysis(temporary_directory, analysis)
            output_dir = Path(temporary_directory) / "output"

            result = self.run_builder(analysis_path, output_dir)

            self.assertEqual(0, result.returncode, result.stderr)
            workbook = load_workbook(next(output_dir.glob("*.xlsx")))
            sheet = workbook["QA-100"]
            self.assertEqual(
                "[POSITIVE — R01] Already prefixed description", sheet["C2"].value
            )
            expected_fills = [PASS_FILL, FAIL_FILL, NOT_RUN_FILL, BLOCKED_FILL, NOT_RUN_FILL]
            actual_fills = [
                sheet[f"G{row}"].fill.fgColor.rgb[-6:] for row in range(2, 7)
            ]
            self.assertEqual(expected_fills, actual_fills)

    def test_repeated_builds_are_semantically_identical(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            first = Path(temporary_directory) / "first"
            second = Path(temporary_directory) / "second"
            first_input = self.write_analysis(first, self.load_valid_analysis())
            second_input = self.write_analysis(second, self.load_valid_analysis())
            first_output = first / "output"
            second_output = second / "output"
            self.assertEqual(0, self.run_builder(first_input, first_output).returncode)
            self.assertEqual(0, self.run_builder(second_input, second_output).returncode)

            self.assertEqual(
                next(first_output.glob("*.html")).read_text(encoding="utf-8"),
                next(second_output.glob("*.html")).read_text(encoding="utf-8"),
            )
            first_doc = Document(next(first_output.glob("*.docx")))
            second_doc = Document(next(second_output.glob("*.docx")))
            self.assertEqual(
                [paragraph.text for paragraph in first_doc.paragraphs],
                [paragraph.text for paragraph in second_doc.paragraphs],
            )
            first_book = load_workbook(next(first_output.glob("*.xlsx")), data_only=False)
            second_book = load_workbook(next(second_output.glob("*.xlsx")), data_only=False)
            for sheet_name in first_book.sheetnames:
                self.assertEqual(
                    list(first_book[sheet_name].values),
                    list(second_book[sheet_name].values),
                )

    def test_invalid_analysis_is_rejected_without_outputs(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            invalid_analysis = json.loads(
                (FIXTURES / "invalid-analysis.json").read_text(encoding="utf-8")
            )
            analysis_path = self.write_analysis(temporary_directory, invalid_analysis)
            output_dir = Path(temporary_directory) / "output"
            result = self.run_builder(analysis_path, output_dir)

            self.assertNotEqual(0, result.returncode)
            self.assertIn("Invalid Netra analysis", result.stderr)
            self.assertFalse(output_dir.exists())


if __name__ == "__main__":
    unittest.main()
